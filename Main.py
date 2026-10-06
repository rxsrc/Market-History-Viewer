"""
Market History Viewer
---------------------
pip install yfinance pandas numpy matplotlib openpyxl gspread

Features
  * Company name next to the ticker (TTWO - Take-Two Interactive Software, Inc.)
  * Search by date (jump to a day) and filter by date range
  * Price chart that follows the date range, with options: range, type
    (line / area / candlestick), SMA 20 / 50, peak & low markers, volume
  * Peak & Low summary for the selected range (or full history if no filter)
  * Export to CSV, Excel (.xlsx) and Google Sheets
  * Multithreaded loading (history + company name fetched in parallel)
  * Split window: compare several stocks side by side (or stacked)
  * Keyboard shortcuts (press F1 in the app for the list)
"""

import os
import queue
import threading
import time
import traceback
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog

import numpy as np
import pandas as pd
import yfinance as yf
import matplotlib.dates as mdates
from matplotlib.figure import Figure
from matplotlib.ticker import FuncFormatter
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

# ---- Theme: royal blue + purple accents, black/grey text, white titles -------
ROYAL = "#4169E1"
PURPLE = "#7B2FBE"
LAV = "#EDE7F6"        # light purple tint (buttons)
LAV2 = "#D9CBEF"       # hover
FOUND = "#CDB4F0"      # row highlighted by "Find Date"
BG = "#F4F5FA"         # main background
SIDE = "#E6E9F4"       # sidebar background
TEXT = "#111111"       # black text
GREY = "#5A5F6B"       # grey text
LINE_GREY = "#C9CCD6"
UP_TINT = "#E3F1E3"    # gain rows (light green tint, black text)
DOWN_TINT = "#F7E1E1"  # loss rows (light red tint, black text)

CACHE_TTL = 300  # seconds
SERVICE_ACCOUNT_FILE = os.environ.get(
    "GOOGLE_SERVICE_ACCOUNT",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "service_account.json"),
)

SHORTCUTS = """\
Ctrl+N        Split: add another stock pane
Ctrl+W        Close the active pane
Ctrl+1..4     Switch active pane
Ctrl+L        Focus the ticker box
Ctrl+G        Focus the 'go to date' box
F5            Refresh the active pane
F6            Toggle side-by-side / stacked
Ctrl+S        Export CSV
Ctrl+E        Export Excel
Ctrl+Shift+G  Export to Google Sheets
F1            This help
Enter         Load ticker / jump to date
"""


# --------------------------------------------------------------------------
# Worker-thread helpers (no Tk calls in here)
# --------------------------------------------------------------------------
def fetch_stock(ticker):
    """Fetch full history and company name in parallel. Returns (name, df)."""
    box = {"name": ""}

    def get_name():
        try:
            info = yf.Ticker(ticker).info
            box["name"] = info.get("longName") or info.get("shortName") or ""
        except Exception:
            box["name"] = ""

    name_thread = threading.Thread(target=get_name, daemon=True)
    name_thread.start()

    df = yf.Ticker(ticker).history(period="max")
    if df.empty:
        raise ValueError(f"No data found for '{ticker}'.")

    if getattr(df.index, "tz", None) is not None:
        df.index = df.index.tz_localize(None)
    df.index = df.index.normalize()
    df = df.sort_index()

    open_ = df["Open"].where(df["Open"] != 0)
    df["Net"] = df["Close"] - df["Open"]
    df["Pct"] = df["Net"] / open_ * 100

    name_thread.join(timeout=10)
    return box["name"], df


def _clean(x):
    if hasattr(x, "item"):
        x = x.item()
    if isinstance(x, float) and pd.isna(x):
        return None
    return x


def write_excel(path, df, ticker, name):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = ticker[:31]
    ws.append(list(df.columns))
    head_fill = PatternFill("solid", fgColor=ROYAL[1:])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    g_fill, r_fill = PatternFill("solid", fgColor=UP_TINT[1:]), PatternFill("solid", fgColor=DOWN_TINT[1:])
    black = Font(color="111111")

    for date, o, c, n, p in df.itertuples(index=False, name=None):
        pct = None if pd.isna(p) else p / 100
        ws.append([datetime.strptime(date, "%Y-%m-%d"), _clean(o), _clean(c), _clean(n), pct])
        r = ws.max_row
        up = not (pd.notna(n) and n < 0)
        for cell in ws[r]:
            cell.fill = g_fill if up else r_fill
            cell.font = black
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 2).number_format = "$#,##0.00"
        ws.cell(r, 3).number_format = "$#,##0.00"
        ws.cell(r, 4).number_format = "+0.00;-0.00;0.00"
        ws.cell(r, 5).number_format = "+0.00%;-0.00%;0.00%"

    for col in "ABCDE":
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_google_sheet(df, title, email, sa_path):
    import gspread

    gc = gspread.service_account(filename=sa_path)
    sh = gc.create(title)
    ws = sh.sheet1
    values = [list(df.columns)] + [
        [_clean(x) for x in row] for row in df.itertuples(index=False, name=None)
    ]
    ws.resize(rows=len(values), cols=len(df.columns))
    ws.update(range_name="A1", values=values)
    if email:
        sh.share(email, perm_type="user", role="writer")
    return sh.url


def parse_date_safe(text):
    text = (text or "").strip()
    if not text:
        return None
    try:
        return pd.to_datetime(text)
    except Exception:
        return None


# --------------------------------------------------------------------------
# One stock pane
# --------------------------------------------------------------------------
class StockPanel(tk.Frame):
    CHUNK = 600
    MAX_CANDLES = 300
    RANGE_CHOICES = ["Auto", "1M", "3M", "6M", "1Y", "5Y", "Max"]
    TYPE_CHOICES = ["Line", "Area", "Candlestick"]
    MONTHS = {"1M": 1, "3M": 3, "6M": 6, "1Y": 12, "5Y": 60}

    def __init__(self, master, app):
        super().__init__(master, bg=BG, highlightthickness=3,
                         highlightbackground=BG, highlightcolor=PURPLE)
        self.app = app
        self.ticker = ""
        self.name = ""
        self.df = None            # full history, ascending
        self.view = None          # date-filtered slice, ascending
        self.filter_active = False
        self.range_text = ("", "")  # (from, to) strings for this pane
        self.request_id = 0
        self._rows, self._pos, self._job = [], 0, None
        self._found = None        # (iid, original_tag) of the "Find Date" row

        # Chart option state
        self.range_var = tk.StringVar(value="Auto")
        self.type_var = tk.StringVar(value="Line")
        self.sma20_var = tk.BooleanVar(value=False)
        self.sma50_var = tk.BooleanVar(value=False)
        self.pk_var = tk.BooleanVar(value=True)
        self.vol_var = tk.BooleanVar(value=False)

        # ---- Title banner (white text on royal blue)
        banner = tk.Frame(self, bg=ROYAL)
        banner.pack(fill=tk.X)
        self.title_lbl = tk.Label(banner, text="Stock Preview", font=("Arial", 16, "bold"),
                                  bg=ROYAL, fg="white", anchor="w")
        self.title_lbl.pack(fill=tk.X, padx=12, pady=8)

        # ---- Info (left) + chart with options (right)
        top = tk.Frame(self, bg=BG)
        top.pack(fill=tk.X, padx=10, pady=(8, 4))
        top.columnconfigure(1, weight=1)

        self.info_lbl = tk.Label(top, text="", justify="left", anchor="nw", bg=BG, fg=TEXT,
                                 font=("Arial", 10), width=24)
        self.info_lbl.grid(row=0, column=0, sticky="nw")

        chart_col = tk.Frame(top, bg=BG)
        chart_col.grid(row=0, column=1, sticky="nsew", padx=(10, 0))

        self.chart_title = tk.Label(chart_col, text="Price chart", font=("Arial", 9, "bold"),
                                    bg=ROYAL, fg="white", anchor="w", padx=8, pady=3)
        self.chart_title.pack(fill=tk.X)

        opts = tk.Frame(chart_col, bg=BG)
        opts.pack(fill=tk.X, pady=(4, 0))
        tk.Label(opts, text="Range", bg=BG, fg=GREY, font=("Arial", 8)).pack(side=tk.LEFT)
        cb_range = ttk.Combobox(opts, textvariable=self.range_var, values=self.RANGE_CHOICES,
                                width=6, state="readonly")
        cb_range.pack(side=tk.LEFT, padx=(3, 10))
        tk.Label(opts, text="Type", bg=BG, fg=GREY, font=("Arial", 8)).pack(side=tk.LEFT)
        cb_type = ttk.Combobox(opts, textvariable=self.type_var, values=self.TYPE_CHOICES,
                               width=11, state="readonly")
        cb_type.pack(side=tk.LEFT, padx=3)
        for cb in (cb_range, cb_type):
            cb.bind("<<ComboboxSelected>>", lambda e: self._draw_chart())

        opts2 = tk.Frame(chart_col, bg=BG)
        opts2.pack(fill=tk.X)
        for text, var in (("SMA 20", self.sma20_var), ("SMA 50", self.sma50_var),
                          ("Peak/Low", self.pk_var), ("Volume", self.vol_var)):
            ttk.Checkbutton(opts2, text=text, variable=var, command=self._draw_chart,
                            style="Chart.TCheckbutton").pack(side=tk.LEFT, padx=(0, 8))

        self.fig = Figure(figsize=(4.6, 2.2), dpi=100, facecolor=BG)
        self.canvas = FigureCanvasTkAgg(self.fig, master=chart_col)
        w = self.canvas.get_tk_widget()
        w.configure(height=210, width=360, highlightthickness=0, bg=BG)
        w.pack(fill=tk.BOTH, expand=True)

        # ---- Peak & Low summary strip
        self.strip_head = tk.Label(self, text="Peak & Low", font=("Arial", 10, "bold"),
                                   bg=PURPLE, fg="white", anchor="w", padx=10, pady=3)
        self.strip_head.pack(fill=tk.X, padx=10)
        cards = tk.Frame(self, bg=BG)
        cards.pack(fill=tk.X, padx=10, pady=(0, 6))
        self.cards = {}
        for i, (key, cap) in enumerate((("peak", "Peak"), ("low", "Low"),
                                        ("spread", "Peak-to-Low Spread"), ("change", "Close Change"))):
            cards.columnconfigure(i, weight=1, uniform="cards")
            card = tk.Frame(cards, bg="white", highlightthickness=1, highlightbackground=PURPLE)
            card.grid(row=0, column=i, sticky="nsew", padx=(0 if i == 0 else 4, 0), pady=(4, 0))
            tk.Label(card, text=cap, bg="white", fg=GREY, font=("Arial", 8)).pack(anchor="w", padx=8, pady=(4, 0))
            val = tk.Label(card, text="-", bg="white", fg=TEXT, font=("Arial", 13, "bold"))
            val.pack(anchor="w", padx=8)
            sub = tk.Label(card, text="", bg="white", fg=GREY, font=("Arial", 8))
            sub.pack(anchor="w", padx=8, pady=(0, 4))
            self.cards[key] = (val, sub)

        # ---- Table
        table_frame = tk.Frame(self, bg=BG)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 10))
        columns = ("Date", "Open", "Close", "Net ($)", "Change (%)")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings")
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=90, minwidth=60, anchor=tk.CENTER)
        self.tree.tag_configure("green", foreground=TEXT, background=UP_TINT)
        self.tree.tag_configure("red", foreground=TEXT, background=DOWN_TINT)
        self.tree.tag_configure("found", foreground=TEXT, background=FOUND)
        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(fill=tk.BOTH, expand=True)

        self._draw_chart()

    # ---- state -----------------------------------------------------------
    def set_active_look(self, active):
        self.configure(highlightbackground=PURPLE if active else BG)

    def set_loading(self, ticker):
        self.ticker = ticker
        self.title_lbl.config(text=f"Fetching data for {ticker}...")

    def set_error(self):
        self.title_lbl.config(text=f"{self.ticker} - not found" if self.ticker else "Stock Preview")

    def show_data(self, ticker, name, df):
        self.ticker, self.name, self.df = ticker, name, df
        self.title_lbl.config(text=f"{ticker} - {name}" if name else ticker)
        start, end = parse_date_safe(self.range_text[0]), parse_date_safe(self.range_text[1])
        self.apply_filter(start, end)

    # ---- filtering / table ----------------------------------------------
    def apply_filter(self, start, end):
        if self.df is None:
            return
        self.filter_active = start is not None or end is not None
        self.view = self.df.loc[start:end]
        self._render_table()
        self._draw_chart()
        self._update_info()
        self._update_stats()

    def _render_table(self):
        self._cancel_job()
        self._found = None
        self.tree.delete(*self.tree.get_children())
        if self.view is None or self.view.empty:
            self._rows = []
            return
        self._rows = list(self.view[["Open", "Close", "Net", "Pct"]].iloc[::-1].itertuples(name=None))
        self._pos = 0
        self._insert_chunk()

    @staticmethod
    def _row_tag(n):
        return "red" if (pd.notna(n) and n < 0) else "green"

    def _insert_chunk(self, limit=None):
        limit = limit or self.CHUNK
        end = min(self._pos + limit, len(self._rows))
        for idx, o, c, n, p in self._rows[self._pos:end]:
            net = "n/a" if pd.isna(n) else f"{'+' if n >= 0 else ''}{n:.2f}"
            pct = "n/a" if pd.isna(p) else f"{'+' if p >= 0 else ''}{p:.2f}%"
            ds = idx.strftime("%Y-%m-%d")
            self.tree.insert("", "end", iid=ds, values=(ds, f"${o:.2f}", f"${c:.2f}", net, pct),
                             tags=(self._row_tag(n),))
        self._pos = end
        self._job = self.after(1, self._insert_chunk) if end < len(self._rows) else None

    def _cancel_job(self):
        if self._job is not None:
            self.after_cancel(self._job)
            self._job = None

    def _flush(self):
        self._cancel_job()
        if self._pos < len(self._rows):
            self._insert_chunk(len(self._rows))

    def _clear_found(self):
        if self._found:
            iid, tag = self._found
            if self.tree.exists(iid):
                self.tree.item(iid, tags=(tag,))
            self._found = None

    def goto_date(self, ts):
        if self.view is None or self.view.empty:
            return "No data loaded in this pane."
        self._flush()
        idx = self.view.index
        pos = max(idx.searchsorted(ts, side="right") - 1, 0)  # last trading day <= ts
        target = idx[pos]
        if target != ts.normalize():
            note = f"No trading on {ts:%Y-%m-%d}; showing nearest day {target:%Y-%m-%d}."
        else:
            note = f"Jumped to {target:%Y-%m-%d}."
        iid = target.strftime("%Y-%m-%d")
        self.update_idletasks()
        self._clear_found()
        self._found = (iid, self._row_tag(self.view.loc[target, "Net"]))
        self.tree.item(iid, tags=("found",))
        self.tree.selection_set(iid)
        self.tree.focus(iid)
        self.tree.see(iid)
        return note

    # ---- info + peak/low summary ----------------------------------------------
    def _update_info(self):
        if self.df is None or self.df.empty:
            self.info_lbl.config(text="")
            return
        last = self.df.iloc[-1]
        self.info_lbl.config(text=(
            f"Last close: ${last['Close']:.2f}\n"
            f"As of: {self.df.index[-1]:%Y-%m-%d}\n"
            f"First trade: {self.df.index[0]:%Y-%m-%d}\n"
            f"Showing {len(self.view):,} of {len(self.df):,} days"
        ))

    def _update_stats(self):
        def blank(msg):
            self.strip_head.config(text=msg)
            for val, sub in self.cards.values():
                val.config(text="-")
                sub.config(text="")

        v = self.view
        if v is None or v.empty:
            blank("Peak & Low - no trading days in that range" if self.df is not None else "Peak & Low")
            return

        highs = v["High"].dropna() if "High" in v else pd.Series(dtype=float)
        lows = v["Low"].dropna() if "Low" in v else pd.Series(dtype=float)
        if highs.empty or lows.empty:
            highs = lows = v["Close"].dropna()
        hi, hi_d = highs.max(), highs.idxmax()
        lo, lo_d = lows.min(), lows.idxmin()
        first, last = v["Close"].iloc[0], v["Close"].iloc[-1]
        chg = (last / first - 1) * 100 if first else float("nan")

        scope = "Selected range" if self.filter_active else "Full history"
        self.strip_head.config(
            text=f"Peak & Low - {scope}: {v.index[0]:%Y-%m-%d} to {v.index[-1]:%Y-%m-%d}  ({len(v):,} days)")
        self.cards["peak"][0].config(text=f"${hi:,.2f}")
        self.cards["peak"][1].config(text=f"intraday high on {hi_d:%Y-%m-%d}")
        self.cards["low"][0].config(text=f"${lo:,.2f}")
        self.cards["low"][1].config(text=f"intraday low on {lo_d:%Y-%m-%d}")
        self.cards["spread"][0].config(text=f"${hi - lo:,.2f}")
        self.cards["spread"][1].config(
            text=f"{(hi / lo - 1) * 100:,.1f}% above the low" if lo > 0 else "")
        self.cards["change"][0].config(text="n/a" if pd.isna(chg) else f"{'+' if chg >= 0 else ''}{chg:.2f}%")
        self.cards["change"][1].config(text=f"${first:,.2f} to ${last:,.2f}")

    # ---- chart -------------------------------------------------------------
    def _chart_data(self):
        """Data to plot: the filtered range in Auto mode, else a window ending at the view's last day."""
        if self.df is None or self.df.empty:
            return None
        choice = self.range_var.get()
        if choice == "Auto":
            if self.filter_active:
                return self.view
            base, months = self.df, 3
        elif choice == "Max":
            return self.view
        else:
            base, months = self.view, self.MONTHS[choice]
        if base is None or base.empty:
            return base
        return base[base.index >= base.index[-1] - pd.DateOffset(months=months)]

    def _placeholder(self, msg):
        self.fig.clear()
        ax = self.fig.add_subplot(111)
        ax.set_facecolor(BG)
        ax.text(0.5, 0.5, msg, ha="center", va="center", color=GREY, transform=ax.transAxes)
        ax.set_xticks([])
        ax.set_yticks([])
        for s in ax.spines.values():
            s.set_visible(False)
        self.chart_title.config(text="Price chart")
        self.canvas.draw_idle()

    def _draw_chart(self):
        data = self._chart_data()
        if data is None:
            self._placeholder("Load a ticker to see the chart")
            return
        if data.empty:
            self._placeholder("No data in that range")
            return

        self.fig.clear()
        show_vol = self.vol_var.get() and "Volume" in data
        if show_vol:
            gs = self.fig.add_gridspec(4, 1, hspace=0.08)
            ax = self.fig.add_subplot(gs[:3, 0])
            axv = self.fig.add_subplot(gs[3, 0], sharex=ax)
        else:
            ax = self.fig.add_subplot(111)
            axv = None

        def style(a):
            a.set_facecolor("white")
            a.tick_params(colors=GREY, labelsize=7)
            a.grid(alpha=0.35, color=LINE_GREY)
            for s in a.spines.values():
                s.set_color(LINE_GREY)
            for s in ("top", "right"):
                a.spines[s].set_visible(False)

        style(ax)
        if axv is not None:
            style(axv)

        x = mdates.date2num(data.index.values)
        close = data["Close"].values.astype(float)
        kind = self.type_var.get()
        note = ""
        if kind == "Candlestick" and len(data) > self.MAX_CANDLES:
            kind, note = "Line", f" (candles need <= {self.MAX_CANDLES} days)"

        want_extremes = kind == "Candlestick" or self.pk_var.get()
        lo_y = float(np.nanmin(data["Low"].values if want_extremes else close))
        hi_y = float(np.nanmax(data["High"].values if want_extremes else close))
        rng = (hi_y - lo_y) or max(abs(hi_y) * 0.02, 1.0)
        y_bottom, y_top = lo_y - 0.12 * rng, hi_y + 0.16 * rng

        marker = "o" if len(data) <= 40 else None
        if kind == "Candlestick":
            op = data["Open"].values.astype(float)
            up = close >= op
            ax.vlines(x, data["Low"].values, data["High"].values, color=GREY, linewidth=0.8)
            height = np.maximum(np.abs(close - op), rng * 0.003)
            ax.bar(x, height, bottom=np.minimum(op, close), width=0.6,
                   color=[ROYAL if u else PURPLE for u in up])
        else:
            ax.plot(x, close, color=ROYAL, linewidth=1.6, marker=marker, markersize=3, label="Close")
            if kind == "Area":
                ax.fill_between(x, close, y_bottom, color=ROYAL, alpha=0.18)

        has_legend = kind != "Candlestick"
        for n, color, var in ((20, PURPLE, self.sma20_var), (50, "#444444", self.sma50_var)):
            if var.get():
                sma = self.df["Close"].rolling(n).mean().reindex(data.index)
                if sma.notna().any():
                    ax.plot(x, sma.values, color=color, linewidth=1.2, label=f"SMA {n}")
                    has_legend = True

        span = (x[-1] - x[0]) or 1.0

        def ha(xv):
            frac = (xv - x[0]) / span
            return "left" if frac < 0.15 else ("right" if frac > 0.85 else "center")

        if self.pk_var.get():
            hi_d, lo_d = data["High"].idxmax(), data["Low"].idxmin()
            hx, lx = mdates.date2num(hi_d.to_pydatetime()), mdates.date2num(lo_d.to_pydatetime())
            hi, lo = data["High"].max(), data["Low"].min()
            ax.axhline(hi, color=PURPLE, linestyle=":", linewidth=0.9, alpha=0.7)
            ax.axhline(lo, color=ROYAL, linestyle=":", linewidth=0.9, alpha=0.7)
            ax.scatter([hx], [hi], s=34, color=PURPLE, edgecolor=TEXT, zorder=5)
            ax.scatter([lx], [lo], s=34, color=ROYAL, edgecolor=TEXT, zorder=5)
            ax.annotate(f"Peak ${hi:,.2f}  {hi_d:%b %d, %Y}", (hx, hi), xytext=(0, 7),
                        textcoords="offset points", ha=ha(hx), fontsize=7, color=TEXT)
            ax.annotate(f"Low ${lo:,.2f}  {lo_d:%b %d, %Y}", (lx, lo), xytext=(0, -13),
                        textcoords="offset points", ha=ha(lx), fontsize=7, color=TEXT)

        ax.set_ylim(y_bottom, y_top)
        pad = max(span * 0.02, 0.6)
        ax.set_xlim(x[0] - pad, x[-1] + pad)
        ax.yaxis.set_major_formatter(
            FuncFormatter(lambda v, _: f"${v:,.0f}" if abs(v) >= 100 else f"${v:,.2f}"))
        if has_legend:
            ax.legend(fontsize=7, loc="upper left", frameon=False, labelcolor=GREY)

        bottom = ax
        if axv is not None:
            vol = data["Volume"].values.astype(float)
            if len(data) <= self.MAX_CANDLES:
                axv.bar(x, vol, width=0.8, color=PURPLE, alpha=0.55)
            else:
                axv.fill_between(x, vol, step="mid", color=PURPLE, alpha=0.55)
            axv.set_yticks([])
            axv.set_ylabel("Vol", fontsize=7, color=GREY)
            ax.tick_params(labelbottom=False)
            bottom = axv

        bottom.xaxis_date()
        loc = mdates.AutoDateLocator(minticks=3, maxticks=6)
        bottom.xaxis.set_major_locator(loc)
        bottom.xaxis.set_major_formatter(mdates.ConciseDateFormatter(loc))

        self.fig.subplots_adjust(left=0.13, right=0.98, top=0.96, bottom=0.12 if show_vol else 0.14)
        self.chart_title.config(
            text=f"{self.ticker}:  {data.index[0]:%b %d, %Y} to {data.index[-1]:%b %d, %Y}"
                 f"  |  {kind}{note}")
        self.canvas.draw_idle()

    # ---- export data ---------------------------------------------------------
    def export_df(self):
        if self.view is None or self.view.empty:
            return None
        v = self.view.iloc[::-1]
        return pd.DataFrame({
            "Date": v.index.strftime("%Y-%m-%d"),
            "Open": v["Open"].round(4).values,
            "Close": v["Close"].round(4).values,
            "Net ($)": v["Net"].round(4).values,
            "Change (%)": v["Pct"].round(4).values,
        })


# --------------------------------------------------------------------------
# Main application
# --------------------------------------------------------------------------
class StockGUI:
    MAX_PANES = 4

    def __init__(self, root):
        self.root = root
        self.root.title("Market History Viewer")
        self.root.geometry("1300x780")
        self.root.minsize(950, 600)
        self.root.configure(bg=BG)

        self.pool = ThreadPoolExecutor(max_workers=8)
        self.ui_q = queue.Queue()
        self.cache = {}  # ticker -> (timestamp, name, df)
        self.panels = []
        self.active = None
        self.preview_active = False

        self._setup_styles()
        self._build_sidebar()
        self.paned = ttk.PanedWindow(root, orient=tk.HORIZONTAL)  # packed on first use
        self.add_panel()

        self._bind_shortcuts()
        self.root.after(100, self._poll_queue)

    # ---- theme -------------------------------------------------------------
    def _setup_styles(self):
        st = ttk.Style()
        st.theme_use("clam")

        st.configure("Side.TButton", background=LAV, foreground=TEXT, bordercolor=PURPLE,
                     lightcolor=LAV, darkcolor=LAV, focuscolor=LAV, padding=(4, 4),
                     font=("Arial", 9, "bold"))
        st.map("Side.TButton", background=[("active", LAV2), ("pressed", LAV2)],
               foreground=[("active", TEXT)])

        st.configure("Chart.TCheckbutton", background=BG, foreground=TEXT, font=("Arial", 8),
                     focuscolor=BG)
        st.map("Chart.TCheckbutton", background=[("active", BG)], foreground=[("active", TEXT)])

        st.configure("TCombobox", foreground=TEXT, background=LAV, arrowcolor=PURPLE,
                     bordercolor=LINE_GREY, selectbackground="white", selectforeground=TEXT)
        st.map("TCombobox", fieldbackground=[("readonly", "white")],
               foreground=[("readonly", TEXT)], selectbackground=[("readonly", "white")],
               selectforeground=[("readonly", TEXT)])

        st.configure("Treeview", rowheight=22, background="white", fieldbackground="white",
                     foreground=TEXT, bordercolor=LINE_GREY)
        st.configure("Treeview.Heading", background=PURPLE, foreground="white",
                     font=("Arial", 9, "bold"), relief="flat")
        st.map("Treeview.Heading", background=[("active", ROYAL)])
        st.map("Treeview", background=[("selected", FOUND)], foreground=[("selected", TEXT)])

        st.configure("Vertical.TScrollbar", background=LAV2, troughcolor=BG, arrowcolor=PURPLE,
                     bordercolor=LINE_GREY)
        st.configure("TPanedwindow", background=LAV2)

    # ---- sidebar -----------------------------------------------------------
    def _build_sidebar(self):
        left = tk.Frame(self.root, width=240, bg=SIDE)
        left.pack(side=tk.LEFT, fill=tk.Y)
        left.pack_propagate(False)
        self.left = left

        tk.Label(left, text="Market History", bg=ROYAL, fg="white",
                 font=("Arial", 15, "bold"), pady=12).pack(fill=tk.X)
        tk.Frame(left, bg=PURPLE, height=4).pack(fill=tk.X)

        def header(text):
            tk.Label(left, text=text, bg=ROYAL, fg="white", font=("Arial", 10, "bold"),
                     anchor="w", padx=10, pady=3).pack(fill=tk.X, pady=(12, 6))

        def make_entry(parent, width=14, font=("Arial", 12)):
            return tk.Entry(parent, font=font, width=width, justify="center", bg="white", fg=TEXT,
                            insertbackground=TEXT, relief="flat", highlightthickness=1,
                            highlightbackground="#B7BCCB", highlightcolor=PURPLE)

        def button_row(items):
            row = tk.Frame(left, bg=SIDE)
            row.pack(padx=10, fill=tk.X, pady=1)
            for i, (text, cmd) in enumerate(items):
                row.columnconfigure(i, weight=1, uniform="b")
                ttk.Button(row, text=text, command=cmd, style="Side.TButton").grid(
                    row=0, column=i, sticky="ew", padx=2)
            return row

        def stack_buttons(rows):
            for items in rows:
                button_row(items)

        header("Ticker")
        self.ticker_entry = make_entry(left, width=12, font=("Arial", 14))
        self.ticker_entry.pack(pady=(0, 6))
        self.ticker_entry.bind("<Return>", lambda e: self.load_preview())
        button_row([("Load Preview", self.load_preview)])

        header("Go to date (YYYY-MM-DD)")
        self.date_entry = make_entry(left, width=14)
        self.date_entry.pack(pady=(0, 6))
        self.date_entry.bind("<Return>", lambda e: self.goto_date())
        button_row([("Find Date", self.goto_date)])

        header("Filter date range")
        fields = tk.Frame(left, bg=SIDE)
        fields.pack(padx=10)
        for col, cap in enumerate(("From", "To")):
            tk.Label(fields, text=cap, bg=SIDE, fg=GREY, font=("Arial", 8)).grid(row=0, column=col)
        self.from_entry = make_entry(fields, width=10, font=("Arial", 10))
        self.to_entry = make_entry(fields, width=10, font=("Arial", 10))
        self.from_entry.grid(row=1, column=0, padx=3, pady=(0, 6))
        self.to_entry.grid(row=1, column=1, padx=3, pady=(0, 6))
        for e in (self.from_entry, self.to_entry):
            e.bind("<Return>", lambda ev: self.apply_range())
        button_row([("Apply", self.apply_range), ("Clear", self.clear_range)])
        tk.Label(left, text="Either box may be left blank.\nThe chart follows the range.",
                 bg=SIDE, fg=GREY, font=("Arial", 8), justify="center").pack(pady=(4, 0))

        header("Export")
        button_row([("CSV", self.export_csv), ("Excel", self.export_excel),
                    ("Sheets", self.export_sheets)])

        header("Panes")
        stack_buttons([
            [("Split (Ctrl+N)", lambda: self.add_panel(focus=True)), ("Close (Ctrl+W)", self.close_active)],
            [("Layout (F6)", self.toggle_layout), ("Help (F1)", self.show_help)],
        ])

        self.status = tk.Label(left, text="", fg=GREY, bg=SIDE, font=("Arial", 9),
                               wraplength=210, justify="center")
        self.status.pack(pady=12)

    # ---- panes -------------------------------------------------------------
    def _reveal(self):
        if not self.preview_active:
            self.paned.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)
            self.preview_active = True

    def add_panel(self, focus=False):
        if len(self.panels) >= self.MAX_PANES:
            self.root.bell()
            self.set_status(f"Maximum of {self.MAX_PANES} panes.")
            return None
        panel = StockPanel(self.paned, self)
        self.paned.add(panel, weight=1)
        self.panels.append(panel)
        self.set_active(panel, sync=True)
        if focus:
            self._reveal()
            self.ticker_entry.focus_set()
        return panel

    def close_active(self):
        if len(self.panels) <= 1:
            self.set_status("At least one pane must stay open.")
            return
        panel = self.active
        idx = self.panels.index(panel)
        self.panels.remove(panel)
        panel.request_id += 1  # discard any in-flight result
        panel._cancel_job()
        self.paned.forget(panel)
        panel.destroy()
        self.set_active(self.panels[max(0, idx - 1)], sync=True)

    def set_active(self, panel, sync=False):
        changed = panel is not self.active
        self.active = panel
        for p in self.panels:
            p.set_active_look(p is panel)
        if changed or sync:
            for entry, text in ((self.ticker_entry, panel.ticker),
                                (self.from_entry, panel.range_text[0]),
                                (self.to_entry, panel.range_text[1])):
                entry.delete(0, tk.END)
                entry.insert(0, text)

    def toggle_layout(self):
        cur = str(self.paned.cget("orient"))
        new = tk.VERTICAL if cur == tk.HORIZONTAL else tk.HORIZONTAL
        try:
            self.paned.configure(orient=new)
        except tk.TclError:
            self.set_status("Layout toggle isn't supported by this Tk version.")

    def _on_click(self, event):
        w = event.widget
        while w is not None:
            if isinstance(w, StockPanel):
                if w is not self.active:
                    self.set_active(w)
                return
            w = getattr(w, "master", None)

    # ---- shortcuts -----------------------------------------------------------
    def _bind_shortcuts(self):
        r = self.root
        r.bind_all("<Button-1>", self._on_click, add="+")
        r.bind_all("<Control-n>", lambda e: self.add_panel(focus=True))
        r.bind_all("<Control-w>", lambda e: self.close_active())
        r.bind_all("<Control-l>", lambda e: (self.ticker_entry.focus_set(),
                                             self.ticker_entry.select_range(0, tk.END)))
        r.bind_all("<Control-g>", lambda e: (self.date_entry.focus_set(),
                                             self.date_entry.select_range(0, tk.END)))
        r.bind_all("<F5>", lambda e: self.load_preview(force=True))
        r.bind_all("<F6>", lambda e: self.toggle_layout())
        r.bind_all("<F1>", lambda e: self.show_help())
        r.bind_all("<Control-s>", lambda e: self.export_csv())
        r.bind_all("<Control-e>", lambda e: self.export_excel())
        r.bind_all("<Control-G>", lambda e: self.export_sheets())
        r.bind_all("<Control-Shift-G>", lambda e: self.export_sheets())
        for i in range(1, self.MAX_PANES + 1):
            r.bind_all(f"<Control-Key-{i}>", lambda e, i=i: self._select_pane(i - 1))

    def _select_pane(self, i):
        if i < len(self.panels):
            self.set_active(self.panels[i])

    def show_help(self):
        messagebox.showinfo("Keyboard shortcuts", SHORTCUTS)

    # ---- threading plumbing --------------------------------------------------
    def run_bg(self, fn, on_ok=None, on_err=None):
        def task():
            try:
                res = fn()
            except Exception as exc:  # noqa: BLE001
                if on_err:
                    self.ui_q.put(lambda exc=exc: on_err(exc))
                return
            if on_ok:
                self.ui_q.put(lambda: on_ok(res))
        self.pool.submit(task)

    def _poll_queue(self):
        try:
            while True:
                self.ui_q.get_nowait()()
        except queue.Empty:
            pass
        except Exception:  # keep the poller alive
            traceback.print_exc()
        self.root.after(100, self._poll_queue)

    def set_status(self, text):
        self.status.config(text=text)

    # ---- loading -----------------------------------------------------------
    def load_preview(self, force=False):
        ticker = self.ticker_entry.get().upper().strip()
        if not ticker:
            if force and self.active and self.active.ticker:
                ticker = self.active.ticker
            else:
                messagebox.showwarning("Error", "Please enter a valid stock symbol.")
                return

        self._reveal()
        panel = self.active
        panel.request_id += 1
        rid = panel.request_id
        panel.range_text = (self.from_entry.get(), self.to_entry.get())  # applied once data arrives
        panel.set_loading(ticker)

        cached = self.cache.get(ticker)
        if cached and not force and time.time() - cached[0] < CACHE_TTL:
            panel.show_data(ticker, cached[1], cached[2])
            return

        def done(result):
            name, df = result
            self.cache[ticker] = (time.time(), name, df)
            if panel.winfo_exists() and panel.request_id == rid:
                panel.show_data(ticker, name, df)

        def failed(exc):
            if panel.winfo_exists() and panel.request_id == rid:
                panel.set_error()
                messagebox.showerror("Error", f"Could not retrieve data for '{ticker}': {exc}")

        self.run_bg(lambda: fetch_stock(ticker), done, failed)

    # ---- date search / filter -----------------------------------------------------
    def goto_date(self):
        if self.active.df is None:
            self.set_status("Load a ticker first.")
            return
        text = self.date_entry.get().strip()
        if not text:
            return
        ts = parse_date_safe(text)
        if ts is None:
            messagebox.showwarning("Date", "Use a date like 2024-03-15.")
            return
        self.set_status(self.active.goto_date(ts))

    def apply_range(self):
        p = self.active
        if p.df is None:
            self.set_status("Load a ticker first.")
            return
        f_text, t_text = self.from_entry.get().strip(), self.to_entry.get().strip()
        start, end = parse_date_safe(f_text), parse_date_safe(t_text)
        if (f_text and start is None) or (t_text and end is None):
            messagebox.showwarning("Date", "Use dates like 2024-03-15.")
            return
        p.range_text = (f_text, t_text)
        p.apply_filter(start, end)
        n = 0 if p.view is None else len(p.view)
        self.set_status(f"Showing {n:,} trading days." if n else "No trading days in that range.")

    def clear_range(self):
        self.from_entry.delete(0, tk.END)
        self.to_entry.delete(0, tk.END)
        p = self.active
        p.range_text = ("", "")
        if p.df is not None:
            p.apply_filter(None, None)
            self.set_status("Filter cleared.")

    # ---- export --------------------------------------------------------------
    def _export_frame(self):
        p = self.active
        df = p.export_df()
        if df is None:
            messagebox.showwarning("Export", "Nothing to export yet - load a ticker first.")
            return None, None
        return p, df

    def export_csv(self):
        p, df = self._export_frame()
        if df is None:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV", "*.csv")],
            initialfile=f"{p.ticker}_history.csv")
        if not path:
            return
        self.set_status("Exporting CSV...")
        self.run_bg(lambda: df.to_csv(path, index=False),
                    lambda _: self.set_status(f"Saved {os.path.basename(path)}"),
                    lambda e: messagebox.showerror("Export failed", str(e)))

    def export_excel(self):
        p, df = self._export_frame()
        if df is None:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"{p.ticker}_history.xlsx")
        if not path:
            return
        self.set_status("Exporting Excel...")
        self.run_bg(lambda: write_excel(path, df, p.ticker, p.name),
                    lambda _: self.set_status(f"Saved {os.path.basename(path)}"),
                    lambda e: messagebox.showerror("Export failed", str(e)))

    def export_sheets(self):
        p, df = self._export_frame()
        if df is None:
            return

        if os.path.exists(SERVICE_ACCOUNT_FILE):
            email = simpledialog.askstring(
                "Google Sheets", "Your Google account email (the sheet will be shared with it):")
            if not email:
                return
            title = f"{p.ticker} history ({datetime.now():%Y-%m-%d})"
            self.set_status("Uploading to Google Sheets...")

            def ok(url):
                self.set_status("Google Sheet created.")
                webbrowser.open(url)

            self.run_bg(lambda: write_google_sheet(df, title, email.strip(), SERVICE_ACCOUNT_FILE),
                        ok, lambda e: messagebox.showerror("Google Sheets failed", str(e)))
        else:
            # No credentials: copy as tab-separated text and open a blank sheet to paste into.
            self.root.clipboard_clear()
            self.root.clipboard_append(df.to_csv(sep="\t", index=False))
            self.root.update()
            webbrowser.open("https://sheets.new")
            messagebox.showinfo(
                "Google Sheets",
                "The data is on your clipboard and a blank Google Sheet is opening.\n"
                "Click cell A1 and press Ctrl+V.\n\n"
                "For one-click uploads, put a Google service-account key at:\n"
                f"{SERVICE_ACCOUNT_FILE}")


if __name__ == "__main__":
    root = tk.Tk()
    app = StockGUI(root)
    root.mainloop()
