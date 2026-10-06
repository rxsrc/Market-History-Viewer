"""
Market History Viewer
---------------------
pip install yfinance pandas matplotlib openpyxl gspread

Features
  * Company name next to the ticker (TTWO - Take-Two Interactive Software, Inc.)
  * Search by date (jump to a day) and filter by date range
  * 3-month price chart at the top right, between the title and the data sheet
  * Export to CSV, Excel (.xlsx, colour-coded) and Google Sheets
  * Multithreaded loading (history + company name fetched in parallel, UI never freezes)
  * Split window: compare several stocks side by side (or stacked)
  * Keyboard shortcuts (press F1 in the app for the list)
"""

import os
import queue
import threading
import time
import traceback
import webbrowser
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog

import pandas as pd
import yfinance as yf
import matplotlib.dates as mdates
from matplotlib.figure import Figure
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

DARK = "#2c3e50"
LIGHT = "#ecf0f1"
ACCENT = "#2980b9"
GREEN_FG, GREEN_BG = "#276A3C", "#E2EFDA"
RED_FG, RED_BG = "#9C0006", "#FCE4D6"

CACHE_TTL = 300  # seconds
SERVICE_ACCOUNT_FILE = os.environ.get(
    "GOOGLE_SERVICE_ACCOUNT",
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "service_account.json"),
)

SHORTCUTS = """\
Ctrl+N        Split: add another stock pane
Ctrl+W        Close the active pane
Ctrl+1..4     Switch active pane
Ctrl+L        Focus the ticker box
Ctrl+G        Focus the 'go to date' box
F5            Refresh the active pane
F6            Toggle side-by-side / stacked
Ctrl+S        Export CSV
Ctrl+E        Export Excel
Ctrl+Shift+G  Export to Google Sheets
F1            This help
Enter         Load ticker / jump to date
"""


# --------------------------------------------------------------------------
# Worker-thread helpers (no Tk calls in here)
# --------------------------------------------------------------------------
def fetch_stock(ticker):
    """Fetch full history and company name in parallel. Returns (name, df)."""
    box = {"name": ""}

    def get_name():
        try:
            info = yf.Ticker(ticker).info
            box["name"] = info.get("longName") or info.get("shortName") or ""
        except Exception:
            box["name"] = ""

    name_thread = threading.Thread(target=get_name, daemon=True)
    name_thread.start()

    df = yf.Ticker(ticker).history(period="max")
    if df.empty:
        raise ValueError(f"No data found for '{ticker}'.")

    if getattr(df.index, "tz", None) is not None:
        df.index = df.index.tz_localize(None)
    df.index = df.index.normalize()
    df = df.sort_index()

    open_ = df["Open"].where(df["Open"] != 0)
    df["Net"] = df["Close"] - df["Open"]
    df["Pct"] = df["Net"] / open_ * 100

    name_thread.join(timeout=10)
    return box["name"], df


def _clean(x):
    if hasattr(x, "item"):
        x = x.item()
    if isinstance(x, float) and pd.isna(x):
        return None
    return x


def write_excel(path, df, ticker, name):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = ticker[:31]
    ws.append(list(df.columns))
    head_fill = PatternFill("solid", fgColor="2C3E50")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")

    g_fill, r_fill = PatternFill("solid", fgColor=GREEN_BG[1:]), PatternFill("solid", fgColor=RED_BG[1:])
    g_font, r_font = Font(color=GREEN_FG[1:]), Font(color=RED_FG[1:])

    for date, o, c, n, p in df.itertuples(index=False, name=None):
        pct = None if pd.isna(p) else p / 100
        ws.append([datetime.strptime(date, "%Y-%m-%d"), _clean(o), _clean(c), _clean(n), pct])
        r = ws.max_row
        up = not (pd.notna(n) and n < 0)
        for cell in ws[r]:
            cell.fill = g_fill if up else r_fill
            cell.font = g_font if up else r_font
        ws.cell(r, 1).number_format = "yyyy-mm-dd"
        ws.cell(r, 2).number_format = "$#,##0.00"
        ws.cell(r, 3).number_format = "$#,##0.00"
        ws.cell(r, 4).number_format = "+0.00;-0.00;0.00"
        ws.cell(r, 5).number_format = "+0.00%;-0.00%;0.00%"

    for col in "ABCDE":
        ws.column_dimensions[col].width = 15
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def write_google_sheet(df, title, email, sa_path):
    import gspread

    gc = gspread.service_account(filename=sa_path)
    sh = gc.create(title)
    ws = sh.sheet1
    values = [list(df.columns)] + [
        [_clean(x) for x in row] for row in df.itertuples(index=False, name=None)
    ]
    ws.resize(rows=len(values), cols=len(df.columns))
    ws.update(range_name="A1", values=values)
    if email:
        sh.share(email, perm_type="user", role="writer")
    return sh.url


# --------------------------------------------------------------------------
# One stock pane
# --------------------------------------------------------------------------
class StockPanel(tk.Frame):
    CHUNK = 600

    def __init__(self, master, app):
        super().__init__(master, bg=LIGHT, highlightthickness=3,
                         highlightbackground=LIGHT, highlightcolor=ACCENT)
        self.app = app
        self.ticker = ""
        self.name = ""
        self.df = None      # full history, ascending
        self.view = None    # date-filtered slice, ascending
        self.request_id = 0
        self._rows, self._pos, self._job = [], 0, None

        # Title
        self.title_lbl = tk.Label(self, text="Stock Preview", font=("Arial", 16, "bold"),
                                  bg=LIGHT, anchor="w")
        self.title_lbl.pack(fill=tk.X, padx=10, pady=(8, 0))

        # Info (left) + 3-month chart (right), below the title and above the table
        top = tk.Frame(self, bg=LIGHT)
        top.pack(fill=tk.X, padx=10, pady=4)
        top.columnconfigure(1, weight=1)

        self.info_lbl = tk.Label(top, text="", justify="left", anchor="nw", bg=LIGHT,
                                 font=("Arial", 10), width=26)
        self.info_lbl.grid(row=0, column=0, sticky="nw")

        self.fig = Figure(figsize=(3.6, 1.7), dpi=100, facecolor=LIGHT)
        self.ax = self.fig.add_subplot(111)
        self.canvas = FigureCanvasTkAgg(self.fig, master=top)
        w = self.canvas.get_tk_widget()
        w.configure(height=170, width=320, highlightthickness=0)
        w.grid(row=0, column=1, sticky="nsew", padx=(10, 0))
        self._draw_chart()

        # Table
        table_frame = tk.Frame(self, bg=LIGHT)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=(4, 10))
        columns = ("Date", "Open", "Close", "Net ($)", "Change (%)")
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings")
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=90, minwidth=60, anchor=tk.CENTER)
        self.tree.tag_configure("green", foreground=GREEN_FG, background=GREEN_BG)
        self.tree.tag_configure("red", foreground=RED_FG, background=RED_BG)
        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(fill=tk.BOTH, expand=True)

    # ---- state -----------------------------------------------------------
    def set_active_look(self, active):
        self.configure(highlightbackground=ACCENT if active else LIGHT)

    def set_loading(self, ticker):
        self.ticker = ticker
        self.title_lbl.config(text=f"Fetching data for {ticker}...")

    def set_error(self):
        self.title_lbl.config(text=f"{self.ticker} - not found" if self.ticker else "Stock Preview")

    def show_data(self, ticker, name, df):
        self.ticker, self.name, self.df = ticker, name, df
        self.title_lbl.config(text=f"{ticker} - {name}" if name else ticker)
        self.apply_filter(None, None)

    # ---- filtering / table ----------------------------------------------
    def apply_filter(self, start, end):
        if self.df is None:
            return
        self.view = self.df.loc[start:end]
        self._render_table()
        self._draw_chart()
        self._update_info()

    def _render_table(self):
        self._cancel_job()
        self.tree.delete(*self.tree.get_children())
        if self.view is None or self.view.empty:
            self._rows = []
            return
        self._rows = list(self.view[["Open", "Close", "Net", "Pct"]].iloc[::-1].itertuples(name=None))
        self._pos = 0
        self._insert_chunk()

    def _insert_chunk(self, limit=None):
        limit = limit or self.CHUNK
        end = min(self._pos + limit, len(self._rows))
        for idx, o, c, n, p in self._rows[self._pos:end]:
            net = "n/a" if pd.isna(n) else f"{'+' if n >= 0 else ''}{n:.2f}"
            pct = "n/a" if pd.isna(p) else f"{'+' if p >= 0 else ''}{p:.2f}%"
            tag = "red" if (pd.notna(n) and n < 0) else "green"
            ds = idx.strftime("%Y-%m-%d")
            self.tree.insert("", "end", iid=ds, values=(ds, f"${o:.2f}", f"${c:.2f}", net, pct),
                             tags=(tag,))
        self._pos = end
        self._job = self.after(1, self._insert_chunk) if end < len(self._rows) else None

    def _cancel_job(self):
        if self._job is not None:
            self.after_cancel(self._job)
            self._job = None

    def _flush(self):
        self._cancel_job()
        if self._pos < len(self._rows):
            self._insert_chunk(len(self._rows))

    def goto_date(self, ts):
        if self.view is None or self.view.empty:
            return "No data loaded in this pane."
        self._flush()
        idx = self.view.index
        pos = idx.searchsorted(ts, side="right") - 1  # last trading day <= ts
        note = ""
        if pos < 0:
            pos = 0
        target = idx[pos]
        if target != ts.normalize():
            note = f"No trading on {ts:%Y-%m-%d}; showing nearest day {target:%Y-%m-%d}."
        else:
            note = f"Jumped to {target:%Y-%m-%d}."
        iid = target.strftime("%Y-%m-%d")
        self.update_idletasks()
        self.tree.selection_set(iid)
        self.tree.focus(iid)
        self.tree.see(iid)
        return note

    def _update_info(self):
        if self.df is None or self.df.empty:
            self.info_lbl.config(text="")
            return
        last = self.df.iloc[-1]
        last_date = self.df.index[-1]
        recent = self.df.loc[self.df.index >= last_date - pd.DateOffset(months=3), "Close"]
        chg = (recent.iloc[-1] / recent.iloc[0] - 1) * 100 if len(recent) > 1 else 0.0
        self.info_lbl.config(text=(
            f"Last close: ${last['Close']:.2f}\n"
            f"As of: {last_date:%Y-%m-%d}\n"
            f"3-month change: {'+' if chg >= 0 else ''}{chg:.2f}%\n"
            f"First trade: {self.df.index[0]:%Y-%m-%d}\n"
            f"Showing {len(self.view):,} of {len(self.df):,} days"
        ))

    # ---- chart -------------------------------------------------------------
    def _draw_chart(self):
        ax = self.ax
        ax.clear()
        ax.set_facecolor(LIGHT)
        if self.df is None or self.df.empty:
            ax.text(0.5, 0.5, "3-month chart", ha="center", va="center",
                    color="#7f8c8d", transform=ax.transAxes)
            ax.set_xticks([])
            ax.set_yticks([])
        else:
            last_date = self.df.index[-1]
            recent = self.df.loc[self.df.index >= last_date - pd.DateOffset(months=3), "Close"]
            up = recent.iloc[-1] >= recent.iloc[0]
            color = "#27ae60" if up else "#c0392b"
            ax.plot(recent.index, recent.values, color=color, linewidth=1.6)
            ax.fill_between(recent.index, recent.values, recent.min(), color=color, alpha=0.15)
            ax.set_title(f"{self.ticker} - last 3 months (Close)", fontsize=9)
            ax.xaxis.set_major_locator(mdates.AutoDateLocator(maxticks=5))
            ax.xaxis.set_major_formatter(mdates.DateFormatter("%b %d"))
            ax.tick_params(labelsize=7)
            ax.grid(alpha=0.25)
            for s in ("top", "right"):
                ax.spines[s].set_visible(False)
        self.fig.tight_layout()
        self.canvas.draw_idle()

    # ---- export data ---------------------------------------------------------
    def export_df(self):
        if self.view is None or self.view.empty:
            return None
        v = self.view.iloc[::-1]
        return pd.DataFrame({
            "Date": v.index.strftime("%Y-%m-%d"),
            "Open": v["Open"].round(4).values,
            "Close": v["Close"].round(4).values,
            "Net ($)": v["Net"].round(4).values,
            "Change (%)": v["Pct"].round(4).values,
        })


# --------------------------------------------------------------------------
# Main application
# --------------------------------------------------------------------------
class StockGUI:
    MAX_PANES = 4

    def __init__(self, root):
        self.root = root
        self.root.title("Market History Viewer")
        self.root.geometry("1250x720")
        self.root.minsize(900, 560)

        self.pool = ThreadPoolExecutor(max_workers=8)
        self.ui_q = queue.Queue()
        self.cache = {}  # ticker -> (timestamp, name, df)
        self.panels = []
        self.active = None
        self.preview_active = False

        style = ttk.Style()
        style.configure("Treeview", rowheight=22)

        self._build_sidebar()
        self.paned = ttk.PanedWindow(root, orient=tk.HORIZONTAL)  # packed on first use
        self.add_panel()

        self._bind_shortcuts()
        self.root.after(100, self._poll_queue)

    # ---- sidebar -----------------------------------------------------------
    def _build_sidebar(self):
        left = tk.Frame(self.root, width=215, bg=DARK)
        left.pack(side=tk.LEFT, fill=tk.Y)
        left.pack_propagate(False)
        self.left = left

        def label(text, pady=(14, 3)):
            tk.Label(left, text=text, fg="white", bg=DARK,
                     font=("Arial", 11, "bold")).pack(pady=pady)

        def button(text, cmd):
            b = tk.Button(left, text=text, font=("Arial", 9, "bold"), command=cmd, width=20)
            b.pack(pady=2)
            return b

        def entry():
            e = tk.Entry(left, font=("Arial", 12), width=12, justify="center")
            e.pack(pady=3)
            return e

        label("Enter Ticker:", pady=(22, 3))
        self.ticker_entry = tk.Entry(left, font=("Arial", 14), width=10, justify="center")
        self.ticker_entry.pack(pady=4)
        self.ticker_entry.bind("<Return>", lambda e: self.load_preview())
        button("Load Preview", self.load_preview)

        label("Go to date (YYYY-MM-DD):")
        self.date_entry = entry()
        self.date_entry.bind("<Return>", lambda e: self.goto_date())
        button("Find Date", self.goto_date)

        label("Filter date range:", pady=(14, 2))
        self.from_entry = entry()
        self.to_entry = entry()
        tk.Label(left, text="(from  /  to; either may be blank)", fg="#bdc3c7",
                 bg=DARK, font=("Arial", 8)).pack()
        row = tk.Frame(left, bg=DARK)
        row.pack(pady=3)
        tk.Button(row, text="Apply", width=9, command=self.apply_range).pack(side=tk.LEFT, padx=2)
        tk.Button(row, text="Clear", width=9, command=self.clear_range).pack(side=tk.LEFT, padx=2)

        label("Export:")
        button("CSV  (Ctrl+S)", self.export_csv)
        button("Excel  (Ctrl+E)", self.export_excel)
        button("Google Sheets  (Ctrl+Shift+G)", self.export_sheets)

        label("Panes:")
        button("Split / Add Stock  (Ctrl+N)", lambda: self.add_panel(focus=True))
        button("Close Pane  (Ctrl+W)", self.close_active)
        button("Side-by-side / Stacked  (F6)", self.toggle_layout)
        button("Shortcuts  (F1)", self.show_help)

        self.status = tk.Label(left, text="", fg="#f1c40f", bg=DARK, font=("Arial", 9),
                               wraplength=190, justify="center")
        self.status.pack(pady=10)

    # ---- panes -------------------------------------------------------------
    def _reveal(self):
        if not self.preview_active:
            self.paned.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)
            self.preview_active = True

    def add_panel(self, focus=False):
        if len(self.panels) >= self.MAX_PANES:
            self.root.bell()
            self.set_status(f"Maximum of {self.MAX_PANES} panes.")
            return None
        panel = StockPanel(self.paned, self)
        self.paned.add(panel, weight=1)
        self.panels.append(panel)
        self.set_active(panel, sync_entry=True)
        if focus:
            self._reveal()
            self.ticker_entry.focus_set()
        return panel

    def close_active(self):
        if len(self.panels) <= 1:
            self.set_status("At least one pane must stay open.")
            return
        panel = self.active
        idx = self.panels.index(panel)
        self.panels.remove(panel)
        panel.request_id += 1  # discard any in-flight result
        panel._cancel_job()
        self.paned.forget(panel)
        panel.destroy()
        self.set_active(self.panels[max(0, idx - 1)], sync_entry=True)

    def set_active(self, panel, sync_entry=False):
        changed = panel is not self.active
        self.active = panel
        for p in self.panels:
            p.set_active_look(p is panel)
        if changed or sync_entry:
            self.ticker_entry.delete(0, tk.END)
            self.ticker_entry.insert(0, panel.ticker)

    def toggle_layout(self):
        cur = str(self.paned.cget("orient"))
        new = tk.VERTICAL if cur == tk.HORIZONTAL else tk.HORIZONTAL
        try:
            self.paned.configure(orient=new)
        except tk.TclError:
            self.set_status("Layout toggle isn't supported by this Tk version.")

    def _on_click(self, event):
        w = event.widget
        while w is not None:
            if isinstance(w, StockPanel):
                if w is not self.active:
                    self.set_active(w)
                return
            w = getattr(w, "master", None)

    # ---- shortcuts -----------------------------------------------------------
    def _bind_shortcuts(self):
        r = self.root
        r.bind_all("<Button-1>", self._on_click, add="+")
        r.bind_all("<Control-n>", lambda e: self.add_panel(focus=True))
        r.bind_all("<Control-w>", lambda e: self.close_active())
        r.bind_all("<Control-l>", lambda e: (self.ticker_entry.focus_set(),
                                             self.ticker_entry.select_range(0, tk.END)))
        r.bind_all("<Control-g>", lambda e: (self.date_entry.focus_set(),
                                             self.date_entry.select_range(0, tk.END)))
        r.bind_all("<F5>", lambda e: self.load_preview(force=True))
        r.bind_all("<F6>", lambda e: self.toggle_layout())
        r.bind_all("<F1>", lambda e: self.show_help())
        r.bind_all("<Control-s>", lambda e: self.export_csv())
        r.bind_all("<Control-e>", lambda e: self.export_excel())
        r.bind_all("<Control-G>", lambda e: self.export_sheets())
        r.bind_all("<Control-Shift-G>", lambda e: self.export_sheets())
        for i in range(1, self.MAX_PANES + 1):
            r.bind_all(f"<Control-Key-{i}>", lambda e, i=i: self._select_pane(i - 1))

    def _select_pane(self, i):
        if i < len(self.panels):
            self.set_active(self.panels[i])

    def show_help(self):
        messagebox.showinfo("Keyboard shortcuts", SHORTCUTS)

    # ---- threading plumbing --------------------------------------------------
    def run_bg(self, fn, on_ok=None, on_err=None):
        def task():
            try:
                res = fn()
            except Exception as exc:  # noqa: BLE001
                if on_err:
                    self.ui_q.put(lambda exc=exc: on_err(exc))
                return
            if on_ok:
                self.ui_q.put(lambda: on_ok(res))
        self.pool.submit(task)

    def _poll_queue(self):
        try:
            while True:
                self.ui_q.get_nowait()()
        except queue.Empty:
            pass
        except Exception:  # keep the poller alive
            traceback.print_exc()
        self.root.after(100, self._poll_queue)

    def set_status(self, text):
        self.status.config(text=text)

    # ---- loading -----------------------------------------------------------
    def load_preview(self, force=False):
        ticker = self.ticker_entry.get().upper().strip()
        if not ticker:
            if force and self.active and self.active.ticker:
                ticker = self.active.ticker
            else:
                messagebox.showwarning("Error", "Please enter a valid stock symbol.")
                return

        self._reveal()
        panel = self.active
        panel.request_id += 1
        rid = panel.request_id
        panel.set_loading(ticker)

        cached = self.cache.get(ticker)
        if cached and not force and time.time() - cached[0] < CACHE_TTL:
            panel.show_data(ticker, cached[1], cached[2])
            return

        def done(result):
            name, df = result
            self.cache[ticker] = (time.time(), name, df)
            if panel.winfo_exists() and panel.request_id == rid:
                panel.show_data(ticker, name, df)

        def failed(exc):
            if panel.winfo_exists() and panel.request_id == rid:
                panel.set_error()
                messagebox.showerror("Error", f"Could not retrieve data for '{ticker}': {exc}")

        self.run_bg(lambda: fetch_stock(ticker), done, failed)

    # ---- date search / filter -----------------------------------------------------
    @staticmethod
    def _parse_date(text):
        text = text.strip()
        if not text:
            return None
        return pd.to_datetime(text)

    def goto_date(self):
        if self.active.df is None:
            self.set_status("Load a ticker first.")
            return
        try:
            ts = self._parse_date(self.date_entry.get())
        except Exception:
            messagebox.showwarning("Date", "Use a date like 2024-03-15.")
            return
        if ts is None:
            return
        self.set_status(self.active.goto_date(ts))

    def apply_range(self):
        if self.active.df is None:
            self.set_status("Load a ticker first.")
            return
        try:
            start = self._parse_date(self.from_entry.get())
            end = self._parse_date(self.to_entry.get())
        except Exception:
            messagebox.showwarning("Date", "Use dates like 2024-03-15.")
            return
        self.active.apply_filter(start, end)
        n = 0 if self.active.view is None else len(self.active.view)
        self.set_status(f"Showing {n:,} trading days.")

    def clear_range(self):
        self.from_entry.delete(0, tk.END)
        self.to_entry.delete(0, tk.END)
        if self.active.df is not None:
            self.active.apply_filter(None, None)
            self.set_status("Filter cleared.")

    # ---- export --------------------------------------------------------------
    def _export_frame(self):
        p = self.active
        df = p.export_df()
        if df is None:
            messagebox.showwarning("Export", "Nothing to export yet - load a ticker first.")
            return None, None
        return p, df

    def export_csv(self):
        p, df = self._export_frame()
        if df is None:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".csv", filetypes=[("CSV", "*.csv")],
            initialfile=f"{p.ticker}_history.csv")
        if not path:
            return
        self.set_status("Exporting CSV...")
        self.run_bg(lambda: df.to_csv(path, index=False),
                    lambda _: self.set_status(f"Saved {os.path.basename(path)}"),
                    lambda e: messagebox.showerror("Export failed", str(e)))

    def export_excel(self):
        p, df = self._export_frame()
        if df is None:
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"{p.ticker}_history.xlsx")
        if not path:
            return
        self.set_status("Exporting Excel...")
        self.run_bg(lambda: write_excel(path, df, p.ticker, p.name),
                    lambda _: self.set_status(f"Saved {os.path.basename(path)}"),
                    lambda e: messagebox.showerror("Export failed", str(e)))

    def export_sheets(self):
        p, df = self._export_frame()
        if df is None:
            return

        if os.path.exists(SERVICE_ACCOUNT_FILE):
            email = simpledialog.askstring(
                "Google Sheets", "Your Google account email (the sheet will be shared with it):")
            if not email:
                return
            title = f"{p.ticker} history ({datetime.now():%Y-%m-%d})"
            self.set_status("Uploading to Google Sheets...")

            def ok(url):
                self.set_status("Google Sheet created.")
                webbrowser.open(url)

            self.run_bg(lambda: write_google_sheet(df, title, email.strip(), SERVICE_ACCOUNT_FILE),
                        ok, lambda e: messagebox.showerror("Google Sheets failed", str(e)))
        else:
            # No credentials: copy as tab-separated text and open a blank sheet to paste into.
            self.root.clipboard_clear()
            self.root.clipboard_append(df.to_csv(sep="\t", index=False))
            self.root.update()
            webbrowser.open("https://sheets.new")
            messagebox.showinfo(
                "Google Sheets",
                "The data is on your clipboard and a blank Google Sheet is opening.\n"
                "Click cell A1 and press Ctrl+V.\n\n"
                "For one-click uploads, put a Google service-account key at:\n"
                f"{SERVICE_ACCOUNT_FILE}")


if __name__ == "__main__":
    root = tk.Tk()
    app = StockGUI(root)
    root.mainloop()
